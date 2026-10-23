"""
articles.py — Article management router for Leila IT Help Desk Agent.
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import datetime, timezone
from typing import List, Optional

import aiofiles
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from sqlalchemy.orm import Session as DBSession

from database import get_db
from models import (
    Article,
    ArticleChunk,
    AuditLog,
    CoverageGapCluster,
    Document,
    DocumentChunk,
)
from schemas import (
    ArticleApproveRequest,
    ArticleCreate,
    ArticleOut,
    ArticleRejectRequest,
    ArticleUpdate,
)
from services.groq_service import get_embedding

router = APIRouter()


def _utcnow() -> datetime:
    return datetime.now(timezone.utc).replace(tzinfo=None)


def _new_uuid() -> str:
    return str(uuid.uuid4())


def log_audit(
    db: DBSession,
    event_type: str,
    actor_id: Optional[str],
    event_details: dict,
    data_classification_tag: str = "Public",
) -> None:
    entry = AuditLog(
        id=_new_uuid(),
        event_type=event_type,
        actor_id=actor_id,
        event_details=json.dumps(event_details),
        data_classification_tag=data_classification_tag,
        timestamp=_utcnow(),
    )
    db.add(entry)
    db.commit()


# ---------------------------------------------------------------------------
# Chunking helper
# ---------------------------------------------------------------------------

def _chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    """Split text into overlapping chunks."""
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start += chunk_size - overlap
    return chunks


def _create_article_chunks(db: DBSession, article: Article) -> None:
    """Generate and store ArticleChunks with embeddings for an article."""
    # Remove existing chunks
    db.query(ArticleChunk).filter(ArticleChunk.article_id == article.id).delete()
    db.commit()

    # Build full text from content + steps
    full_text = article.content or ""
    if article.steps:
        try:
            steps = json.loads(article.steps)
            full_text += "\n\nTroubleshooting Steps:\n" + "\n".join(
                f"{i+1}. {s}" for i, s in enumerate(steps)
            )
        except Exception:
            pass

    # Add title at the front for better retrieval
    if article.title:
        full_text = f"{article.title}\n\n{full_text}"

    raw_chunks = _chunk_text(full_text, chunk_size=500, overlap=50)
    approved = article.lifecycle_status == "published"

    for order, chunk_text in enumerate(raw_chunks):
        emb = get_embedding(chunk_text)
        chunk = ArticleChunk(
            id=_new_uuid(),
            article_id=article.id,
            chunk_text=chunk_text,
            embedding_json=json.dumps(emb),
            chunk_order=order,
            approved_status=approved,
            created_at=_utcnow(),
        )
        db.add(chunk)
    db.commit()


# ---------------------------------------------------------------------------
# GET /articles
# ---------------------------------------------------------------------------

@router.get("", response_model=List[ArticleOut])
def list_articles(
    lifecycle_status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    db: DBSession = Depends(get_db),
):
    query = db.query(Article)
    if lifecycle_status:
        query = query.filter(Article.lifecycle_status == lifecycle_status)
    if search:
        query = query.filter(
            (Article.title.ilike(f"%{search}%")) | (Article.content.ilike(f"%{search}%"))
        )
    articles = query.order_by(Article.updated_at.desc()).all()
    result = []
    for a in articles:
        out = _article_to_schema(a)
        result.append(out)
    return result


# ---------------------------------------------------------------------------
# GET /articles/coverage-gaps
# ---------------------------------------------------------------------------

@router.get("/coverage-gaps")
def list_coverage_gaps(db: DBSession = Depends(get_db)):
    gaps = (
        db.query(CoverageGapCluster)
        .order_by(CoverageGapCluster.occurrence_count.desc())
        .all()
    )
    result = []
    for g in gaps:
        result.append(
            {
                "id": g.id,
                "representative_query": g.representative_query,
                "occurrence_count": g.occurrence_count,
                "member_query_ids": json.loads(g.member_query_ids or "[]"),
                "classification_tag": g.classification_tag,
                "suggestion_status": g.suggestion_status,
                "created_at": g.created_at,
                "updated_at": g.updated_at,
            }
        )
    return result


# ---------------------------------------------------------------------------
# GET /articles/{id}
# ---------------------------------------------------------------------------

@router.get("/{article_id}", response_model=ArticleOut)
def get_article(article_id: str, db: DBSession = Depends(get_db)):
    article = db.query(Article).filter(Article.id == article_id).first()
    if not article:
        raise HTTPException(status_code=404, detail="Article not found")
    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# POST /articles
# ---------------------------------------------------------------------------

@router.post("", response_model=ArticleOut, status_code=201)
def create_article(payload: ArticleCreate, db: DBSession = Depends(get_db)):
    article = Article(
        id=_new_uuid(),
        title=payload.title,
        content=payload.content,
        steps=json.dumps(payload.steps) if payload.steps else None,
        category=payload.category,
        context_tags=json.dumps(payload.context_tags) if payload.context_tags else None,
        author_id=payload.author_id,
        author_type=payload.author_type,
        lifecycle_status="draft",
        version=1,
        created_at=_utcnow(),
        updated_at=_utcnow(),
    )
    db.add(article)
    db.commit()
    db.refresh(article)

    _create_article_chunks(db, article)

    log_audit(
        db,
        "article_created",
        payload.author_id,
        {"article_id": article.id, "title": article.title},
    )

    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# PUT /articles/{id}
# ---------------------------------------------------------------------------

@router.put("/{article_id}", response_model=ArticleOut)
def update_article(
    article_id: str, payload: ArticleUpdate, db: DBSession = Depends(get_db)
):
    article = db.query(Article).filter(Article.id == article_id).first()
    if not article:
        raise HTTPException(status_code=404, detail="Article not found")

    if payload.title is not None:
        article.title = payload.title
    if payload.content is not None:
        article.content = payload.content
    if payload.steps is not None:
        article.steps = json.dumps(payload.steps)
    if payload.category is not None:
        article.category = payload.category
    if payload.context_tags is not None:
        article.context_tags = json.dumps(payload.context_tags)

    article.version = (article.version or 1) + 1
    article.lifecycle_status = "pending_review"
    article.updated_at = _utcnow()
    db.commit()
    db.refresh(article)

    _create_article_chunks(db, article)

    log_audit(
        db,
        "article_updated",
        None,
        {"article_id": article.id, "version": article.version},
    )

    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# POST /articles/{id}/submit
# ---------------------------------------------------------------------------

@router.post("/{article_id}/submit", response_model=ArticleOut)
def submit_article(article_id: str, db: DBSession = Depends(get_db)):
    article = db.query(Article).filter(Article.id == article_id).first()
    if not article:
        raise HTTPException(status_code=404, detail="Article not found")
    if article.lifecycle_status not in ("draft",):
        raise HTTPException(
            status_code=400,
            detail=f"Cannot submit article with status '{article.lifecycle_status}'",
        )
    article.lifecycle_status = "pending_review"
    article.updated_at = _utcnow()
    db.commit()

    log_audit(
        db,
        "article_submitted_for_review",
        article.author_id,
        {"article_id": article.id},
    )

    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# POST /articles/{id}/approve
# ---------------------------------------------------------------------------

@router.post("/{article_id}/approve", response_model=ArticleOut)
def approve_article(
    article_id: str, payload: ArticleApproveRequest, db: DBSession = Depends(get_db)
):
    article = db.query(Article).filter(Article.id == article_id).first()
    if not article:
        raise HTTPException(status_code=404, detail="Article not found")

    # Block self-approval
    if article.author_id == payload.reviewer_id:
        raise HTTPException(
            status_code=403,
            detail="Self-approval is not permitted. A different reviewer must approve this article.",
        )

    if article.lifecycle_status != "pending_review":
        raise HTTPException(
            status_code=400,
            detail=f"Article must be in 'pending_review' status to approve. Current status: '{article.lifecycle_status}'",
        )

    article.lifecycle_status = "published"
    article.reviewer_id = payload.reviewer_id
    article.last_reviewed_date = _utcnow()
    article.updated_at = _utcnow()
    db.commit()
    db.refresh(article)

    # Re-chunk with approved_status=True
    _create_article_chunks(db, article)

    log_audit(
        db,
        "article_approved",
        payload.reviewer_id,
        {"article_id": article.id, "title": article.title},
    )

    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# POST /articles/{id}/reject
# ---------------------------------------------------------------------------

@router.post("/{article_id}/reject", response_model=ArticleOut)
def reject_article(
    article_id: str, payload: ArticleRejectRequest, db: DBSession = Depends(get_db)
):
    article = db.query(Article).filter(Article.id == article_id).first()
    if not article:
        raise HTTPException(status_code=404, detail="Article not found")

    if article.lifecycle_status != "pending_review":
        raise HTTPException(
            status_code=400,
            detail="Article must be in 'pending_review' status to reject.",
        )

    article.lifecycle_status = "draft"
    article.updated_at = _utcnow()
    db.commit()

    log_audit(
        db,
        "article_rejected",
        payload.reviewer_id,
        {"article_id": article.id, "reason": payload.reason},
    )

    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# POST /articles/{id}/retire
# ---------------------------------------------------------------------------

@router.post("/{article_id}/retire", response_model=ArticleOut)
def retire_article(article_id: str, db: DBSession = Depends(get_db)):
    article = db.query(Article).filter(Article.id == article_id).first()
    if not article:
        raise HTTPException(status_code=404, detail="Article not found")

    article.lifecycle_status = "retired"
    article.updated_at = _utcnow()
    db.commit()

    # Disable all chunks
    db.query(ArticleChunk).filter(ArticleChunk.article_id == article_id).update(
        {"approved_status": False}
    )
    db.commit()

    log_audit(
        db,
        "article_retired",
        None,
        {"article_id": article.id},
    )

    return _article_to_schema(article)


# ---------------------------------------------------------------------------
# POST /articles/upload-document
# ---------------------------------------------------------------------------

@router.post("/upload-document")
async def upload_document(
    uploader_id: str = Form(...),
    file: UploadFile = File(...),
    db: DBSession = Depends(get_db),
):
    """
    Accept PDF/DOCX/TXT/MD file, extract text, chunk and embed,
    store as DocumentChunks with approved_status=False.
    """
    filename = file.filename or "unknown"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else "txt"

    if ext not in ("pdf", "docx", "md", "txt"):
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: {ext}. Supported: pdf, docx, md, txt",
        )

    content_bytes = await file.read()
    extracted_text = _extract_text(content_bytes, ext)

    if not extracted_text.strip():
        raise HTTPException(status_code=400, detail="Could not extract text from file")

    # Create Document record
    doc = Document(
        id=_new_uuid(),
        filename=filename,
        source_type=ext,
        uploader_id=uploader_id,
        created_at=_utcnow(),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    # Chunk and embed
    raw_chunks = _chunk_text(extracted_text, chunk_size=500, overlap=50)
    for order, chunk_text in enumerate(raw_chunks):
        emb = get_embedding(chunk_text)
        dc = DocumentChunk(
            id=_new_uuid(),
            document_id=doc.id,
            chunk_text=chunk_text,
            embedding_json=json.dumps(emb),
            chunk_order=order,
            approved_status=False,
            created_at=_utcnow(),
        )
        db.add(dc)
    db.commit()

    log_audit(
        db,
        "document_uploaded",
        uploader_id,
        {"document_id": doc.id, "filename": filename, "chunks": len(raw_chunks)},
    )

    return {
        "document_id": doc.id,
        "filename": filename,
        "chunks_created": len(raw_chunks),
        "approved_status": False,
        "message": "Document uploaded and chunked. Awaiting IT Content Owner review.",
    }


# ---------------------------------------------------------------------------
# Text extraction helpers
# ---------------------------------------------------------------------------

def _extract_text(content_bytes: bytes, ext: str) -> str:
    """Extract plain text from various file formats."""
    if ext == "pdf":
        return _extract_pdf(content_bytes)
    if ext == "docx":
        return _extract_docx(content_bytes)
    # txt / md — decode as UTF-8
    try:
        return content_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return content_bytes.decode("latin-1", errors="replace")


def _extract_pdf(content_bytes: bytes) -> str:
    try:
        import PyPDF2

        reader = PyPDF2.PdfReader(io.BytesIO(content_bytes))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)
    except Exception as exc:
        return f"[PDF extraction error: {exc}]"


def _extract_docx(content_bytes: bytes) -> str:
    try:
        import docx

        doc = docx.Document(io.BytesIO(content_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)
    except Exception as exc:
        return f"[DOCX extraction error: {exc}]"


# ---------------------------------------------------------------------------
# Serialisation helper
# ---------------------------------------------------------------------------

def _article_to_schema(article: Article) -> ArticleOut:
    steps = None
    if article.steps:
        try:
            steps = json.loads(article.steps)
        except Exception:
            steps = None

    context_tags = None
    if article.context_tags:
        try:
            context_tags = json.loads(article.context_tags)
        except Exception:
            context_tags = None

    return ArticleOut(
        id=article.id,
        title=article.title,
        content=article.content,
        steps=steps,
        category=article.category,
        context_tags=context_tags,
        author_id=article.author_id,
        author_type=article.author_type,
        reviewer_id=article.reviewer_id,
        lifecycle_status=article.lifecycle_status,
        review_cycle_days=article.review_cycle_days or 90,
        last_reviewed_date=article.last_reviewed_date,
        version=article.version or 1,
        created_at=article.created_at,
        updated_at=article.updated_at,
    )
